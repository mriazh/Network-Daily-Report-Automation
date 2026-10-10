#!/usr/bin/env python3
"""Generate the standardized daily report Excel template.

Output : config/Template-Daily-Report.xlsx
Sheets : Master + 31 daily sheets ("01" .. "31"), all with identical structure.

Run   : python scripts/generate_template.py
Layout, palette and row/column rules follow docs/design.md.
"""
from __future__ import annotations

import json
import re
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = ROOT / "config"
OUTPUT = CONFIG_DIR / "Template-Daily-Report.xlsx"
WAG_DIR = Path(r"C:\Users\adima\Downloads\Daily-WAG")

REPORT_TITLE = (
    "NETWORK INFRASTRUCTURE \u2013 DAILY SERVICE, "
    "CONNECTIVITY & PERFORMANCE MONITORING REPORT"
)

# --- palette (docs/design.md section 2) ---------------------------------------
NAVY = "FF1B365D"
SLATE = "FF2C4D75"
HEADER_BLUE = "FF4A6984"
WARM = "FFF0F4F8"
ICE = "FFEBF2F7"
ZEBRA = "FFF8FAFC"
BORDER_C = "FFD0D7DE"
TEXT = "FF1F2933"
MUTED = "FF52606D"
WHITE = "FFFFFFFF"
GREEN_F, GREEN_T = "FFD4EDDA", "FF155724"
YELLOW_F, YELLOW_T = "FFFFF3CD", "FF856404"
RED_F, RED_T = "FFF8D7DA", "FF721C24"

FONT_NAME = "Calibri"
LAST_COL = 8  # columns A..H
MIN_WIDTHS = [6, 32, 26, 16, 16, 38, 15, 15]  # docs/design.md section 4

F_TITLE = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
F_SECTION = Font(name=FONT_NAME, size=11, bold=True, color=WHITE)
F_SUB = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
F_HEAD = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
F_DATA = Font(name=FONT_NAME, size=10, color=TEXT)
F_TOTAL = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
F_NOTE = Font(name=FONT_NAME, size=9, italic=True, color=MUTED)

THIN = Side(style="thin", color=BORDER_C)
DOUBLE = Side(style="double", color=SLATE)
BORDER_GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_TOTAL = Border(left=THIN, right=THIN, top=DOUBLE, bottom=DOUBLE)
BORDER_SUB = Border(bottom=Side(style="thin", color=HEADER_BLUE))

A_LEFT = Alignment(horizontal="left", vertical="center")
A_CENTER = Alignment(horizontal="center", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")
A_HEAD = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGNMENTS = {"l": A_LEFT, "c": A_CENTER, "r": A_RIGHT, "w": A_WRAP}

# --- reference data loader (prefers config/circuits.json, falls back to config/circuits.example.json) ---
def _load_circuits():
    circuits_file = CONFIG_DIR / "circuits.json"
    if not circuits_file.exists():
        circuits_file = CONFIG_DIR / "circuits.example.json"
    if circuits_file.exists():
        with open(circuits_file, "r", encoding="utf-8") as f:
            data = json.load(f)
            return (
                data.get("pic_names", []),
                [tuple(x) for x in data.get("astinet", [])],
                [tuple(x) for x in data.get("metro_e", [])],
                [tuple(x) for x in data.get("vpn_ip_osa", [])],
                [tuple(x) for x in data.get("telkom_indibiz", [])],
                [tuple(x) for x in data.get("nexthop", [])],
                data.get("ssids", []),
                [(x[0], x[1]) for x in data.get("modems", [])],
            )
    return ([], [], [], [], [], [], [], [])

(
    PIC_NAMES,
    ASTINET_CIRCUITS,
    METRO_E,
    VPN_IP_OSA,
    TELKOM_INDIBIZ,
    NEXTHOP,
    SSIDS,
    MODEMS,
) = _load_circuits()

HOTSPOT_LINKS = [x[0] for x in ASTINET_CIRCUITS if "DPS" in x[0] or "SUB" in x[0]] + \
                [x[0] for x in METRO_E if "T1C" in x[0]] + \
                [x[0] for x in VPN_IP_OSA if "BPN" in x[0]]

# --- helpers -----------------------------------------------------------------
def solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


STATUS_FILLS = {
    "ok": (solid(GREEN_F), Font(name=FONT_NAME, size=10, bold=True, color=GREEN_T)),
    "warn": (solid(YELLOW_F), Font(name=FONT_NAME, size=10, bold=True, color=YELLOW_T)),
    "crit": (solid(RED_F), Font(name=FONT_NAME, size=10, bold=True, color=RED_T)),
}
STATUS_KEYWORDS = (
    ("crit", ("fault", "down", "expired", "critical", "gagal", "rusak")),
    ("warn", ("warning", "high traffic", "monitoring", "paused", "open", ">", "tinggi")),
    ("ok", ("normal", "online", "ok", "active", "aktif", "connected")),
)


def status_style(value: object):
    """Map a status label to its pill fill, or None when unknown/empty."""
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.lower()
    for kind, keywords in STATUS_KEYWORDS:
        if any(keyword in text for keyword in keywords):
            return STATUS_FILLS[kind]
    return None


def text_width(value: object) -> float:
    """Rough width in Excel units for 10pt Calibri (avoids ### and clipping)."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return 9.0
    text = str(value)
    if not text or text.startswith("="):  # formulas render as their result, not source
        return 0.0
    return len(text) * 1.05 + 2.0


FREEZE_AT = "A7"  # banner occupies rows 1-6 on every sheet


def apply_view(worksheet) -> None:
    """Sheet view + print setup. Re-applied to copied sheets: openpyxl's
    copy_worksheet() drops sheet_view, print titles and footers."""
    worksheet.freeze_panes = FREEZE_AT
    worksheet.sheet_view.showGridLines = True
    worksheet.print_title_rows = "1:5"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = worksheet.page_margins.right = 0.4
    worksheet.page_margins.top = worksheet.page_margins.bottom = 0.5
    worksheet.oddFooter.left.text = "Template-Daily-Report"
    worksheet.oddFooter.right.text = "Page &P of &N"


class SheetBuilder:
    """Writes the report layout into one worksheet, tracking row cursor + widths."""

    def __init__(self, worksheet) -> None:
        self.ws = worksheet
        self.row = 1
        self._needs: list[tuple[int, int, object]] = []

    # -- low level -----------------------------------------------------------
    def _measure(self, col: int, span: int, value: object) -> None:
        self._needs.append((col, span, value))

    def _write(
        self,
        row: int,
        col: int,
        span: int,
        value: object = None,
        *,
        font: Font | None = None,
        fill: PatternFill | None = None,
        border: Border | None = None,
        align: Alignment | None = None,
        number_format: str | None = None,
        height: float | None = None,
    ) -> None:
        ws = self.ws
        if span > 1:
            ws.merge_cells(
                start_row=row, start_column=col, end_row=row, end_column=col + span - 1
            )
        for index in range(span):
            cell = ws.cell(row=row, column=col + index)
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if index == 0:
                cell.value = value
                if font is not None:
                    cell.font = font
                if align is not None:
                    cell.alignment = align
                if number_format is not None:
                    cell.number_format = number_format
        self._measure(col, span, value)
        if height is not None:
            ws.row_dimensions[row].height = height

    def spacer(self, rows: int = 1) -> None:
        self.row += rows

    # -- blocks --------------------------------------------------------------
    def banner(self, date_value: str) -> None:
        self.spacer(2)
        self._write(self.row, 1, LAST_COL, REPORT_TITLE, font=F_TITLE, fill=solid(NAVY),
                    align=A_CENTER, height=30)
        self.row += 1
        self._write(self.row, 1, 2, "Date / Period", font=F_SUB, fill=solid(ICE),
                    border=BORDER_GRID, align=A_LEFT)
        self._write(self.row, 3, 2, date_value, font=F_DATA, border=BORDER_GRID, align=A_CENTER)
        self.row += 1
        self._write(self.row, 1, 2, "PIC Network Support", font=F_SUB, fill=solid(ICE),
                    border=BORDER_GRID, align=A_LEFT)
        self._write(self.row, 3, LAST_COL - 2, "  |  ".join(
            f"{i}. {name}" for i, name in enumerate(PIC_NAMES, 1)),
            font=F_DATA, border=BORDER_GRID, align=A_LEFT, height=18)
        self.row += 2

    def section(self, letter: str, emoji: str, title: str) -> None:
        self._write(self.row, 1, LAST_COL, f"{letter}. {emoji} {title}", font=F_SECTION,
                    fill=solid(SLATE), align=A_LEFT, height=22)
        self.row += 1

    def subsection(self, number: str, title: str) -> None:
        self._write(self.row, 1, LAST_COL, f"{number}. {title}", font=F_SUB,
                    fill=solid(WARM), border=BORDER_SUB, align=A_LEFT, height=18)
        self.row += 1

    def note(self, text: str) -> None:
        self._write(self.row, 1, LAST_COL, text, font=F_NOTE, fill=solid(ZEBRA),
                    border=BORDER_GRID, align=A_WRAP, height=16)
        self.row += 1

    # -- tables --------------------------------------------------------------
    def table(
        self,
        columns: Sequence[tuple],
        rows: Sequence[Sequence[object]],
        total_row: Sequence[object] | None = None,
        row_height: float | None = None,
    ) -> None:
        """columns: (header, span, align, number_format); rows: cell values.

        Trailing slack is absorbed by the last column, so short tables still fill A..H.
        `total_row` is a sequence, or a callable(first_data_row, last_data_row).
        """
        columns = list(columns)
        used = sum(span for _h, span, _a, _f in columns)
        if used > LAST_COL:
            raise ValueError(f"tabel melebihi {LAST_COL} kolom: {used}")
        if used < LAST_COL:
            header, span, align, fmt = columns[-1]
            columns[-1] = (header, span + LAST_COL - used, align, fmt)

        start = self.row
        col = 1
        for header, span, _align, _fmt in columns:
            self._write(start, col, span, header, font=F_HEAD, fill=solid(HEADER_BLUE),
                        border=BORDER_GRID, align=A_HEAD, height=24)
            col += span
        self.row += 1

        for index, values in enumerate(rows):
            row = self.row
            col = 1
            stripe = solid(ZEBRA) if index % 2 else None
            for value, (_h, span, align, fmt) in zip(values, columns):
                pill = status_style(value) if align == "c" else None
                fill, font = (pill[0], pill[1]) if pill else (stripe, F_DATA)
                self._write(row, col, span, value, font=font, fill=fill,
                            border=BORDER_GRID, align=ALIGNMENTS[align],
                            number_format=fmt, height=row_height)
                col += span
            self.row += 1

        if total_row is not None:
            values = total_row(start + 1, self.row - 1) if callable(total_row) else total_row
            row = self.row
            col = 1
            for value, (_h, span, align, fmt) in zip(values, columns):
                self._write(row, col, span, value, font=F_TOTAL, fill=solid(WARM),
                            border=BORDER_TOTAL, align=ALIGNMENTS[align],
                            number_format=fmt, height=18)
                col += span
            self.row += 1

    def legend(self) -> None:
        self.note("Note: MRTG values in Mbps; percentage = utilization against link "
                  "capacity. Status: Normal / High Traffic / Fault.")
        self.spacer(1)
        col = 1
        for label, key in (("Normal / Online / OK", "ok"),
                           ("Warning / Monitoring", "warn"),
                           ("Critical / Fault / Expired", "crit")):
            fill, font = STATUS_FILLS[key]
            self._write(self.row, col, 2, label, font=font, fill=fill,
                        border=BORDER_GRID, align=A_CENTER, height=16)
            col += 2
        self.row += 2

    def footer(self) -> None:
        self._write(self.row, 1, LAST_COL,
                    "Daily Report - Network Support |Template Excel | "
                    "scripts/generate_template.py", font=F_NOTE, align=A_CENTER)
        self.row += 1

    # -- finishing -----------------------------------------------------------
    def finalize(self) -> None:
        """Widen columns so no value is truncated (### in numeric, cut text)."""
        widths = list(MIN_WIDTHS)
        for col, span, value in self._needs:
            need = text_width(value)
            block = sum(widths[col - 1:col - 1 + span])
            if span == 1:
                widths[col - 1] = max(widths[col - 1], need)
            elif need > block:
                # Grow the merged block evenly instead of inflating its first column.
                share = (need - block) / span
                for index in range(col - 1, col - 1 + span):
                    widths[index] += share
        for index, width in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = round(min(width, 48), 2)
        apply_view(self.ws)


PCT = '0.00"%"'
GB = "#,##0.00"


def section_a(b: SheetBuilder) -> None:
    b.section("A", "\U0001F310", "Telkomsel IoT M2M")

    b.subsection("1", "Service Status Overview")
    b.table(
        [("No", 1, "c", None), ("Parameter", 2, "l", None), ("Value", 1, "r", None),
         ("Status", 1, "c", None), ("Remarks", 3, "l", None)],
        [
            [1, "Status Service", None, "Normal", None],
            [2, "Total Usage", None, None, None],
            [3, "Active Session", None, None, None],
        ],
    )
    b.note("Usage values (GB) are taken from the Telkomsel IoT M2M portal; "
           "Status is filled with Normal / Fault / Maintenance.")

    b.subsection("2", "Top Daily Usage (Top 5)")
    b.table(
        [("Rank", 1, "c", None), ("Site", 2, "l", None), ("Usage (GB)", 1, "r", GB),
         ("Period", 1, "c", None), ("Remarks", 3, "l", None)],
        [[rank, None, None, None, None] for rank in range(1, 6)],
    )

    b.subsection("3", "Top Usage by Site - Periodic Monitoring")
    b.table(
        [("Rank", 1, "c", None), ("Site", 2, "l", None), ("Usage (GB)", 1, "r", GB),
         ("Monitoring Period", 1, "c", None), ("Remarks", 3, "l", None)],
        [[rank, None, None, None, None] for rank in range(1, 6)],
    )
    b.spacer(1)


def _issue_total(first: int, last: int) -> list:
    """TOTAL row: only links with issue (High Traffic / Fault) are counted."""
    issues = f'(COUNTIF(E{first}:E{last},"High Traffic")+COUNTIF(E{first}:E{last},"Fault"))'
    return [
        "TOTAL", "Links with Issue", None, None, f"={issues}",
        f'=IFERROR((SUMIF(E{first}:E{last},"High Traffic",F{first}:F{last})'
        f'+SUMIF(E{first}:E{last},"Fault",F{first}:F{last}))/{issues},"")',
        None,
    ]


def _mrtg_table(b: SheetBuilder, number: str, title: str, links: Sequence[tuple]) -> None:
    b.subsection(number, title)
    b.table(
        [("No", 1, "c", None), ("Link / Name", 1, "l", None), ("SID", 1, "l", None),
         ("Capacity", 1, "c", None), ("Status", 1, "c", None),
         ("Utilization (%)", 1, "r", PCT), ("Remarks", 2, "l", None)],
        [[i, name, sid, capacity, None, None, None]
         for i, (name, sid, capacity) in enumerate(links, 1)],
        total_row=_issue_total,
    )
    b.spacer(1)


def section_b(b: SheetBuilder) -> None:
    b.section("B", "\U0001F4CA", "Usage by MRTG (Network Connectivity)")
    _mrtg_table(b, "1", "ASTINET", ASTINET_CIRCUITS)
    _mrtg_table(b, "2", "Metro-E", METRO_E)
    _mrtg_table(b, "3", "VPN IP (OSA)", VPN_IP_OSA)
    _mrtg_table(b, "4", "Telkom Indibiz", TELKOM_INDIBIZ)
    b.note("Capacity follows MRTG data (1 GB / 2 MB / 4 MB / 10 MB / 12 MB / 50 MB / "
           "150 MB); Utilization = Usage / Capacity x 100%. The TOTAL row counts only "
           "links with issue (High Traffic / Fault).")
    b.spacer(1)


def section_c(b: SheetBuilder) -> None:
    b.section("C", "\U0001F500", "Switch & Core Infrastructure")

    b.subsection("1", "Nexthop Monitoring (Ping)")
    b.table(
        [("No", 1, "c", None), ("Destination", 1, "l", None), ("Role / Interface", 1, "l", None),
         ("IP Address", 1, "c", None), ("Average (ms)", 1, "r", None),
         ("Loss (%)", 1, "r", PCT), ("Status", 1, "c", None), ("Remarks", 1, "l", None)],
        [[i, name, None, ip, None, None, None, None]
         for i, (name, ip) in enumerate(NEXTHOP, 1)],
    )
    b.note("Reference threshold: Loss < 1% and average < 100 ms means Normal.")
    b.spacer(1)

    b.subsection("2", "PRTG Sensor Monitoring")
    b.table(
        [("No", 1, "c", None), ("Category", 2, "l", None), ("Count", 1, "r", None),
         ("Status", 1, "c", None), ("Remarks", 3, "l", None)],
        [
            [1, "UP", None, None, None],
            [2, "Paused", None, None, None],
        ],
        total_row=lambda first, last: [
            None, "TOTAL SENSORS", f"=SUM(D{first}:D{last})", None, None],
    )
    b.spacer(1)

    b.subsection("3", "Top 10 CPU Utilization")
    b.table(
        [("Rank", 1, "c", None), ("Device", 2, "l", None), ("Utilization (%)", 1, "r", PCT),
         ("Baseline", 1, "r", PCT), ("Status", 1, "c", None), ("Remarks", 2, "l", None)],
        [[i, None, None, 80.0, None, None] for i in range(1, 11)],
    )
    b.note("CPU baseline: Normal < 60%, Warning 60% - 80%, Critical > 80%.")
    b.spacer(1)

    b.subsection("4", "Top 10 Memory Usage")
    b.table(
        [("Rank", 1, "c", None), ("Device", 2, "l", None), ("Utilization (%)", 1, "r", PCT),
         ("Baseline", 1, "r", PCT), ("Status", 1, "c", None), ("Remarks", 2, "l", None)],
        [[i, None, None, 70.0, None, None] for i in range(1, 11)],
    )
    b.note("Memory baseline: Normal < 70%, Monitoring 70% - 85%, Critical > 85%.")
    b.spacer(1)

    b.subsection("5", "Avg Uptime by Device")
    b.table(
        [("No", 1, "c", None), ("Parameter / Scope", 2, "l", None), ("SNMP (%)", 1, "r", PCT),
         ("ICMP (%)", 1, "r", PCT), ("HTTPS (%)", 1, "r", PCT), ("Remarks", 2, "l", None)],
        [[1, "Device Uptime", None, None, None, "Device Uptime xx%"]],
    )
    b.note("HTTPS is filled with '-' when the device has no HTTPS sensor.")
    b.spacer(1)

    b.subsection("6", "Technical Utilization Analysis")
    b.table(
        [("No", 1, "c", None), ("Aspect", 2, "l", None), ("Analysis / Findings", 5, "l", None)],
        [[1, "CPU Utilization", None], [2, "Memory Usage", None], [3, "Uptime & Connectivity", None]],
        row_height=30,
    )
    b.spacer(1)


def section_d(b: SheetBuilder) -> None:
    b.section("D", "\U0001F4F6", "Huawei AP (WAC) & Wi-Fi SSID")

    b.subsection("1", "Access Point Status")
    head = b.row
    b.table(
        [("No", 1, "c", None), ("Parameter", 1, "l", None), ("Description", 2, "l", None),
         ("Count", 1, "r", None), ("Percentage (%)", 1, "r", PCT), ("Status", 2, "c", None)],
        [
            [1, "Normal", "AP online", None,
             f'=IFERROR(E{head + 1}/(E{head + 1}+E{head + 2})*100,"")', None],
            [2, "Fault", "AP offline / alarm", None,
             f'=IFERROR(E{head + 2}/(E{head + 1}+E{head + 2})*100,"")', None],
            [3, "Connected Users", "Total connected users", None, None, None],
        ],
    )
    b.note("Percentage = Count / (Normal + Fault) x 100%.")
    b.spacer(1)

    b.subsection("2", "AP Physical Status")
    b.table(
        [("No", 1, "c", None), ("Category", 1, "l", None), ("Type", 1, "l", None),
         ("OK Count", 1, "r", None), ("Fault Count", 1, "r", None),
         ("Remarks", 2, "l", None), ("Status", 1, "c", None)],
        [
            [1, "Indoor", "Indoor AP", None, None, None, None],
            [2, "Outdoor", "Outdoor AP", None, None, None, None],
        ],
        total_row=lambda first, last: [
            None, "TOTAL", "Indoor + Outdoor", f"=SUM(D{first}:D{last})",
            f"=SUM(E{first}:E{last})", None, None],
    )
    b.spacer(1)

    b.subsection("3", "User Distribution by SSID")
    head = b.row
    total_row_idx = head + 1 + len(SSIDS)
    b.table(
        [("No", 1, "c", None), ("SSID", 3, "l", None), ("Connected Users", 2, "r", None),
         ("Percentage (%)", 2, "r", PCT)],
        [[i, ssid, None,
          f'=IFERROR(E{head + i}/$E${total_row_idx}*100,"")']
         for i, ssid in enumerate(SSIDS, 1)],
        total_row=lambda first, last: [
            None, "TOTAL USERS", f"=SUM(E{first}:E{last})", 100],
    )
    b.note("SSID percentage = Connected Users / Total Users x 100%.")
    b.spacer(1)


def section_e(b: SheetBuilder) -> None:
    b.section("E", "\U0001F4F1", "Modem Orbit")

    rows: list[list[object]] = []
    for number, (location, packages) in enumerate(MODEMS, 1):
        for index, package in enumerate(packages):
            rows.append([number if index == 0 else None, location if index == 0 else "",
                         package, None, None, None, None])

    b.table(
        [("No", 1, "c", None), ("Location / Modem", 1, "l", None), ("Package", 1, "l", None),
         ("Remaining Quota (GB)", 1, "r", GB), ("Total Quota (GB)", 1, "r", GB),
         ("Valid Until", 1, "c", None), ("Status", 2, "c", None)],
        rows,
        total_row=lambda first, last: [
            None, "TOTAL QUOTA", None, f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})",
            None, None],
    )
    b.note("Valid Until and Status (Active / Expired Soon / Expired) are taken "
           "manually from the Orbit application; Expired is highlighted in red.")
    b.spacer(1)


def section_f(b: SheetBuilder) -> None:
    b.section("F", "\u26A0\uFE0F", "Highlights & Incident Follow-up")

    b.subsection("1", "High Traffic")
    b.table(
        [("No", 1, "c", None), ("Link / Item", 3, "l", None), ("Next Action", 3, "l", None), ("Status", 1, "c", None)],
        [[i, None, None, None] for i in range(1, 5)],
    )
    b.spacer(1)

    b.subsection("2", "Hardware Fault")
    b.table(
        [("No", 1, "c", None), ("Device", 3, "l", None), ("Next Action", 3, "l", None), ("Status", 1, "c", None)],
        [[i, None, None, None] for i in range(1, 4)],
    )
    b.spacer(1)

    b.subsection("3", "Problem Open & Follow-up")
    b.table(
        [("No", 1, "c", None), ("Problem Description", 4, "l", None),
         ("Next Action", 2, "l", None), ("Status", 1, "c", None)],
        [[1, None, None, None]],
        row_height=26,
    )
    b.spacer(1)


SECTION_BUILDERS = (section_a, section_b, section_c, section_d, section_e, section_f)


def build_sheet(workbook: Workbook, title: str, date_value: str) -> None:
    """Create one worksheet with the full A-F layout."""
    worksheet = workbook.create_sheet(title)
    b = SheetBuilder(worksheet)
    b.banner(date_value)
    for build in SECTION_BUILDERS:
        build(b)
    b.legend()
    b.footer()
    b.finalize()


def generate(output: Path = OUTPUT) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    build_sheet(workbook, "Master", "________ / ________ / ________")
    master = workbook["Master"]
    master.sheet_properties.tabColor = NAVY

    for day in range(1, 32):
        name = f"{day:02d}"
        worksheet = workbook.copy_worksheet(master)
        worksheet.title = name
        worksheet.sheet_properties.tabColor = SLATE
        apply_view(worksheet)  # copy_worksheet() drops view/print settings
        wag_file = WAG_DIR / f"{day}.txt"
        if wag_file.is_file():
            text = wag_file.read_text(encoding="utf-8", errors="replace")
            date_match = re.search(r"\b\d{1,2}\s+[A-Za-z]+\s+\d{4}\b", text)
            worksheet["C4"] = date_match.group(0) if date_match else f"{day:02d} Oktober 2026"
        elif day <= 11:
            worksheet["C4"] = f"{day:02d} Oktober 2026"
        else:
            worksheet["C4"] = f"{day:02d} / ________ / ________"

    workbook.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output

    workbook.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def verify(path: Path, expected_days: int = 31) -> None:
    """Fail loudly if the generated template is not usable."""
    workbook = load_workbook(path)
    expected_sheets = ["Master"] + [f"{day:02d}" for day in range(1, expected_days + 1)]
    actual = workbook.sheetnames
    if actual != expected_sheets:
        raise AssertionError(f"daftar sheet tidak sesuai: {actual}")

    reference: tuple | None = None
    for name in expected_sheets:
        worksheet = workbook[name]
        values = [cell.value for row in worksheet.iter_rows() for cell in row
                  if cell.value is not None]
        if not values:
            raise AssertionError(f"sheet {name} kosong")
        for required in ("A.", "B.", "C.", "D.", "E.", "F."):
            if not any(isinstance(v, str) and v.startswith(f"{required} ") for v in values):
                raise AssertionError(f"section {required} tidak ada di sheet {name}")
        widths = tuple(round(worksheet.column_dimensions[letter].width or 0, 2)
                       for letter in "ABCDEFGH")
        if max(widths) == 0:
            raise AssertionError(f"lebar kolom sheet {name} belum diset")
        if worksheet.max_column != LAST_COL:
            raise AssertionError(f"sheet {name} lebar kolom {worksheet.max_column}, "
                                 f"harus {LAST_COL}")
        if worksheet.freeze_panes is None:
            raise AssertionError(f"sheet {name} belum freeze panes")
        signature = (worksheet.max_row, len(worksheet.merged_cells.ranges), widths,
                     worksheet.freeze_panes, worksheet.print_title_rows)
        if reference is None:
            reference = signature
        elif signature != reference:
            raise AssertionError(f"struktur sheet {name} berbeda dari Master: "
                                 f"{signature} != {reference}")
        if name != "Master" and not str(worksheet["C4"].value).startswith(name):
            raise AssertionError(f"sheet {name} tanggal di C4 salah: {worksheet['C4'].value!r}")
    workbook.close()
    print(f"OK  {path}")
    print(f"    sheets   : {len(expected_sheets)} (Master + {expected_days} hari)")
    print(f"    baris    : {reference[0]} per sheet, {reference[1]} merge")
    print(f"    kolom    : {LAST_COL} (A-H), freeze {reference[3]}, "
          f"judul cetak {reference[4]}")


def main() -> int:
    path = generate()
    verify(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())